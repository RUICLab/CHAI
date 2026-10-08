"""Build an editable PowerPoint version of the CHAI architecture diagram."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "CHAI_System_Architecture.pptx"


COLORS = {
    "navy": "183B56",
    "ink": "243746",
    "muted": "5D6D7E",
    "panel": "F7F9FB",
    "panel_line": "91A4B6",
    "input_fill": "FCE8E6",
    "input_line": "B9473E",
    "blue_fill": "EAF2F8",
    "blue_line": "39729B",
    "cyan_fill": "DDF4F7",
    "cyan_line": "16879A",
    "orange_fill": "FFF2D8",
    "orange_line": "C97B1A",
    "green_fill": "E5F4E9",
    "green_line": "2F7D4A",
    "purple_fill": "EFE9F8",
    "purple_line": "76539A",
    "white": "FFFFFF",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_text(shape, text, size=12, bold=False, color="ink", align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.07)
    tf.margin_right = Inches(0.07)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(COLORS[color])
    return shape


def add_text(slide, x, y, w, h, text, size=12, bold=False, color="ink",
             align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_text(shape, text, size, bold, color, align)


def add_box(slide, x, y, w, h, text, fill, line, size=11, bold=False,
            radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(COLORS[fill])
    shape.line.color.rgb = rgb(COLORS[line])
    shape.line.width = Pt(1.5)
    set_text(shape, text, size=size, bold=bold)
    return shape


def add_panel(slide, x, y, w, h, title, subtitle=None):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = rgb(COLORS["panel"])
    panel.fill.transparency = 4
    panel.line.color.rgb = rgb(COLORS["panel_line"])
    panel.line.width = Pt(1.3)
    add_text(slide, x + 0.2, y + 0.10, w - 0.4, 0.34, title,
             size=15, bold=True, color="navy", align=PP_ALIGN.LEFT)
    if subtitle:
        add_text(slide, x + 0.2, y + 0.42, w - 0.4, 0.24, subtitle,
                 size=8.5, color="muted", align=PP_ALIGN.LEFT)
    return panel


def _add_arrowhead(connector):
    ln = connector._element.spPr.get_or_add_ln()
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    ln.append(tail)


def add_arrow(slide, x1, y1, x2, y2, dashed=False, color="ink", width=1.7):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    connector.line.color.rgb = rgb(COLORS[color])
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    _add_arrowhead(connector)
    return connector


def add_line(slide, x1, y1, x2, y2, dashed=False, color="ink", width=1.7):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    connector.line.color.rgb = rgb(COLORS[color])
    connector.line.width = Pt(width)
    if dashed:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return connector


def add_stage_label(slide, x, y, w, text):
    return add_text(slide, x, y, w, 0.24, text.upper(), size=7.5,
                    bold=True, color="muted")


def build_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(COLORS["white"])

    add_text(slide, 0.42, 0.16, 12.45, 0.46,
             "CHAI System Architecture", size=24, bold=True,
             color="navy", align=PP_ALIGN.LEFT)
    add_text(slide, 0.44, 0.62, 12.3, 0.27,
             "Constraint-guided offline synthesis and energy-aware runtime dispatch",
             size=10.5, color="muted", align=PP_ALIGN.LEFT)

    # Two architectural regions.
    add_panel(
        slide, 0.35, 1.02, 8.55, 5.92,
        "A. Offline CHAI control plane",
        "Generate, validate, tune, and audit approximation configurations",
    )
    add_panel(
        slide, 9.08, 1.02, 3.90, 5.92,
        "B. Runtime dispatcher",
        "Select an execution level from measured stored energy",
    )

    # Offline top row.
    xs = [0.62, 2.25, 3.95, 5.63, 7.30]
    ws = [1.38, 1.45, 1.45, 1.45, 1.32]
    add_stage_label(slide, xs[0], 1.72, ws[0], "Inputs")
    add_stage_label(slide, xs[1], 1.72, ws[1], "Operator system")
    add_stage_label(slide, xs[2], 1.72, ws[2], "Control plane")
    add_stage_label(slide, xs[3], 1.72, ws[3], "Program analysis")
    add_stage_label(slide, xs[4], 1.72, ws[4], "LLM")

    add_box(slide, xs[0], 2.00, ws[0], 1.16,
            "Source code\nHardware profile\nTechnique cards",
            "input_fill", "input_line", size=10, bold=True)
    add_box(slide, xs[1], 2.00, ws[1], 1.16,
            "Card generation\nSchema validation\nRegistry + prompts",
            "blue_fill", "blue_line", size=9.8)
    add_box(slide, xs[2], 2.00, ws[2], 1.16,
            "App/HW filtering\nCost ranking\nFeasible catalog",
            "blue_fill", "blue_line", size=9.8)
    add_box(slide, xs[3], 2.00, ws[3], 1.16,
            "Function extraction\nCall graph\nPer-function narrowing",
            "blue_fill", "blue_line", size=9.5)
    add_box(slide, xs[4], 2.00, ws[4], 1.16,
            "Plan generation\nCode generation",
            "cyan_fill", "cyan_line", size=10, bold=True)

    for i in range(4):
        add_arrow(slide, xs[i] + ws[i] + 0.03, 2.58, xs[i + 1] - 0.04, 2.58)

    # Offline lower validation/evaluation path. The return route from the LLM
    # is drawn above the boxes so the second row reads naturally left-to-right.
    add_line(slide, 7.96, 3.17, 7.96, 3.48)
    add_line(slide, 7.96, 3.48, 1.60, 3.48)
    add_arrow(slide, 1.60, 3.48, 1.60, 3.78)
    add_stage_label(slide, 0.62, 3.50, 2.05, "Validation and evaluation")
    add_box(slide, 0.62, 3.82, 1.96, 1.28,
            "Plan contract\nOperator conformance\nKnob safety\nClaim pruning/clamping",
            "orange_fill", "orange_line", size=8.9, bold=True)
    add_box(slide, 2.92, 3.82, 1.85, 1.28,
            "Compile\nBehavior + identity\nError validation",
            "orange_fill", "orange_line", size=9.4)
    add_box(slide, 5.10, 3.82, 1.68, 1.28,
            "Fused evaluation\nBayesian knob tuning",
            "purple_fill", "purple_line", size=9.5, bold=True)
    add_box(slide, 7.12, 3.82, 1.52, 1.28,
            "Validated configs\nEnergy + error\nKnob settings",
            "green_fill", "green_line", size=9.4, bold=True)

    add_arrow(slide, 2.61, 4.46, 2.89, 4.46)
    add_arrow(slide, 4.80, 4.46, 5.07, 4.46)
    add_arrow(slide, 6.81, 4.46, 7.09, 4.46)

    # Audit artifacts as a cross-cutting result.
    audit = add_box(slide, 0.62, 5.55, 8.27, 0.72,
                    "Audit artifacts: rejection logs  •  constraint proofs  •  raw responses  •  accepted code  •  decision traces",
                    "blue_fill", "blue_line", size=10.2)
    audit.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_text(slide, 0.80, 6.35, 7.9, 0.28,
             "Static checks establish declared-contract compliance; compilation and measurement establish executable behavior.",
             size=8.4, color="muted", align=PP_ALIGN.LEFT)

    # Runtime panel.
    add_stage_label(slide, 9.38, 1.72, 1.48, "Characterization")
    add_stage_label(slide, 11.06, 1.72, 1.58, "Device state")
    add_box(slide, 9.34, 2.00, 1.56, 1.02,
            "Level table\nL0 / L1 / L2\nEnergy + error",
            "green_fill", "green_line", size=9.4, bold=True)
    add_box(slide, 11.08, 2.00, 1.56, 1.02,
            "Capacitor voltage\nC and Vrelease",
            "input_fill", "input_line", size=9.4, bold=True)

    add_box(slide, 9.77, 3.43, 2.45, 0.76,
            "Available-energy estimate\n+ 10% safety margin",
            "blue_fill", "blue_line", size=10, bold=True)
    add_arrow(slide, 10.12, 3.04, 10.38, 3.40)
    add_arrow(slide, 11.86, 3.04, 11.61, 3.40)

    add_box(slide, 9.61, 4.58, 2.78, 0.82,
            "Select least-approximate\naffordable level",
            "cyan_fill", "cyan_line", size=10.5, bold=True)
    add_arrow(slide, 11.00, 4.21, 11.00, 4.55)

    out_y = 5.83
    out_w = 0.72
    out_xs = [9.40, 10.23, 11.06, 11.89]
    out_labels = ["L0", "L1", "L2", "DEFER"]
    out_fills = ["green_fill", "green_fill", "orange_fill", "input_fill"]
    out_lines = ["green_line", "green_line", "orange_line", "input_line"]
    for x, label, fill, line in zip(out_xs, out_labels, out_fills, out_lines):
        add_box(slide, x, out_y, out_w, 0.48, label, fill, line,
                size=9.2, bold=True)
    add_arrow(slide, 11.00, 5.42, 11.00, 5.74)

    add_text(slide, 9.42, 6.38, 3.16, 0.32,
             "Evaluated with canonical levels in Fused/MSP430 simulation",
             size=8.5, color="muted")

    # Dashed integration path from offline validated configurations to table.
    add_arrow(slide, 8.67, 4.20, 9.31, 2.54, dashed=True,
              color="purple_line", width=1.5)
    label = add_text(slide, 8.40, 3.20, 1.18, 0.42,
                     "Integration\npath",
                     size=7.8, bold=True, color="purple_line")
    label.fill.solid()
    label.fill.fore_color.rgb = rgb(COLORS["white"])
    label.fill.transparency = 6

    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    build_slide(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
